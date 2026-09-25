import csv
import json
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.utils import get_column_letter


# ============================================================
# PATHS
# ============================================================

BASE_DIR = Path(__file__).resolve().parent.parent

if len(sys.argv) > 1:
    WORK_DIR = Path(sys.argv[1])
    if not WORK_DIR.is_absolute():
        WORK_DIR = BASE_DIR / WORK_DIR
else:
    WORK_DIR = BASE_DIR / "output"

WORK_DIR.mkdir(parents=True, exist_ok=True)

INPUT_FILE = WORK_DIR / "final_invoice_data.json"
VALIDATION_FILE = WORK_DIR / "validation_report.json"
CSV_OUTPUT = WORK_DIR / "invoice_output.csv"
EXCEL_OUTPUT = WORK_DIR / "invoice_output.xlsx"


# ============================================================
# HELPERS
# ============================================================

def as_dict(value):
    return value if isinstance(value, dict) else {}


def as_list(value):
    return value if isinstance(value, list) else []


def first_present(mapping, *keys):
    mapping = as_dict(mapping)
    for key in keys:
        value = mapping.get(key)
        if value not in (None, "", []):
            return value
    return None


def document_sections(document):
    return {
        "invoice": as_dict(document.get("invoice")),
        "order": as_dict(document.get("order")),
        "seller": as_dict(document.get("seller")),
        "billing": as_dict(document.get("billing")),
        "shipping": as_dict(document.get("shipping")),
        "supply_delivery": as_dict(document.get("supply_delivery")),
        "totals": as_dict(document.get("totals")),
        "payment": as_dict(document.get("payment")),
    }


def normalize_document_list(data):
    """Return the canonical page-aware documents list.

    The integrator's multi-document format is preferred. A legacy
    single-document structure is supported for compatibility.
    """
    documents = as_list(data.get("documents"))
    if documents:
        return documents

    legacy_keys = (
        "invoice",
        "order",
        "seller",
        "billing",
        "shipping",
        "supply_delivery",
        "items",
        "charges",
        "totals",
        "amount_in_words",
        "reverse_charge",
        "payment",
    )

    if any(key in data for key in legacy_keys):
        return [{
            "page": first_present(data, "page", "source_page"),
            "invoice": as_dict(data.get("invoice")),
            "order": as_dict(data.get("order")),
            "seller": as_dict(data.get("seller")),
            "billing": as_dict(data.get("billing")),
            "shipping": as_dict(data.get("shipping")),
            "supply_delivery": as_dict(data.get("supply_delivery")),
            "items": as_list(data.get("items")),
            "charges": as_list(data.get("charges")),
            "table_other_rows": as_list(data.get("table_other_rows")),
            "totals": as_dict(data.get("totals")),
            "amount_in_words": data.get("amount_in_words", ""),
            "reverse_charge": data.get("reverse_charge", ""),
            "payment": as_dict(data.get("payment")),
        }]

    return []


def row_value(row, *keys):
    return first_present(row, *keys)


def build_invoice_context(document, data):
    sections = document_sections(document)
    invoice = sections["invoice"]
    order = sections["order"]
    seller = sections["seller"]
    billing = sections["billing"]
    shipping = sections["shipping"]
    supply = sections["supply_delivery"]
    totals = sections["totals"]
    payment = sections["payment"]

    return {
        "document_type": data.get("processing", {}).get("document_type", "INVOICE"),
        "page": first_present(document, "page", "source_page"),
        "invoice_number": first_present(invoice, "invoice_number", "number"),
        "invoice_details": first_present(invoice, "invoice_details", "details"),
        "invoice_date": first_present(invoice, "invoice_date", "date"),
        "order_number": first_present(order, "order_number", "number"),
        "order_date": first_present(order, "order_date", "date"),
        "seller_name": first_present(seller, "name", "seller_name"),
        "seller_address": first_present(seller, "address"),
        "seller_pan": first_present(seller, "pan"),
        "seller_gstin": first_present(seller, "gstin", "gst_registration_number"),
        "billing_name": first_present(billing, "name"),
        "billing_address": first_present(billing, "address"),
        "billing_state_code": first_present(billing, "state_code"),
        "shipping_name": first_present(shipping, "name"),
        "shipping_address": first_present(shipping, "address"),
        "shipping_state_code": first_present(shipping, "state_code"),
        "place_of_supply": first_present(supply, "place_of_supply"),
        "place_of_delivery": first_present(supply, "place_of_delivery"),
        "invoice_tax_total": first_present(totals, "tax_amount", "total_tax", "tax_total"),
        "invoice_total": first_present(
            totals,
            "total_amount",
            "resolved_total_amount",
            "document_total",
            "total",
        ),
        "observed_total_amount": first_present(
            totals,
            "observed_total_amount",
        ),
        "total_resolution": first_present(
            totals,
            "total_resolution",
        ),
        "amount_in_words": first_present(
            document,
            "amount_in_words",
        ) or first_present(totals, "amount_in_words"),
        "reverse_charge": first_present(document, "reverse_charge"),
        "transaction_id": first_present(payment, "transaction_id"),
        "payment_date_time": first_present(payment, "date_time", "payment_date_time"),
        "invoice_value": first_present(payment, "invoice_value"),
        "payment_mode": first_present(payment, "mode", "payment_mode"),
    }


def get_products(document):
    return as_list(document.get("items"))


def get_charges(document):
    return as_list(document.get("charges"))


def get_other_rows(document):
    return as_list(document.get("table_other_rows"))


def make_item_row(context, item, row_type):
    return {
        **context,
        "row_type": row_type,
        "charge_type": row_value(item, "charge_type"),
        "description": row_value(item, "description", "product_description", "name"),
        "quantity": row_value(item, "quantity", "qty"),
        "unit_price": row_value(item, "unit_price", "unit_amount", "price", "amount"),
        "discount": row_value(item, "discount"),
        "net_amount": row_value(item, "net_amount", "net"),
        "tax_rate": row_value(item, "tax_rate"),
        "tax_type": row_value(item, "tax_type"),
        "tax_amount": row_value(item, "tax_amount", "tax"),
        "total_amount": row_value(item, "total_amount", "total", "amount_total"),
        "asin": row_value(item, "asin"),
        "hsn": row_value(item, "hsn", "hsn_code"),
        "validation": row_value(item, "validation"),
    }


def make_document_row(context):
    return {
        **context,
        "row_type": "document",
        "charge_type": None,
        "description": None,
        "quantity": None,
        "unit_price": None,
        "discount": None,
        "net_amount": None,
        "tax_rate": None,
        "tax_type": None,
        "tax_amount": None,
        "total_amount": None,
        "asin": None,
        "hsn": None,
        "validation": None,
    }


# ============================================================
# LOAD INPUT
# ============================================================

print("=" * 70)
print("UNIVERSAL INVOICE CSV + EXCEL EXPORT")
print("=" * 70)
print(f"Working directory : {WORK_DIR}")
print(f"Input file        : {INPUT_FILE}")

if not INPUT_FILE.exists():
    raise FileNotFoundError(
        f"\nFinal invoice JSON not found:\n{INPUT_FILE}\n\n"
        "Please run invoice_integrator.py first."
    )

with open(INPUT_FILE, "r", encoding="utf-8") as handle:
    data = json.load(handle)

validation = None
if VALIDATION_FILE.exists():
    with open(VALIDATION_FILE, "r", encoding="utf-8") as handle:
        validation = json.load(handle)
    print("Validation report loaded.")
else:
    print("Validation report not found; export will continue without it.")

if not isinstance(data, dict):
    raise ValueError("final_invoice_data.json must contain a JSON object.")

documents = normalize_document_list(data)
if not documents:
    raise ValueError(
        "No invoice documents were found in final_invoice_data.json."
    )

print(f"Documents detected : {len(documents)}")


# ============================================================
# BUILD FLAT EXPORT ROWS
# ============================================================

contexts = []
all_rows = []

for index, document in enumerate(documents, start=1):
    if not isinstance(document, dict):
        continue

    context = build_invoice_context(document, data)
    context["document_index"] = index
    contexts.append((document, context))

    products = get_products(document)
    charges = get_charges(document)
    other_rows = get_other_rows(document)

    if products:
        for item in products:
            all_rows.append(make_item_row(context, item, "product"))

    if charges:
        for charge in charges:
            all_rows.append(make_item_row(context, charge, "charge"))

    # Preserve rows that the table extractor did not classify as a
    # product/charge instead of silently dropping observations.
    for other in other_rows:
        all_rows.append(make_item_row(context, other, "other"))

    # A charge-only invoice must still be represented even when there
    # are no product rows. Likewise, an empty/partially extracted
    # document remains visible at document level.
    if not products and not charges and not other_rows:
        all_rows.append(make_document_row(context))


# ============================================================
# CSV EXPORT
# ============================================================

csv_fields = [
    "document_index",
    "document_type",
    "page",
    "invoice_number",
    "invoice_details",
    "invoice_date",
    "order_number",
    "order_date",
    "seller_name",
    "seller_address",
    "seller_pan",
    "seller_gstin",
    "billing_name",
    "billing_address",
    "billing_state_code",
    "shipping_name",
    "shipping_address",
    "shipping_state_code",
    "place_of_supply",
    "place_of_delivery",
    "row_type",
    "charge_type",
    "description",
    "quantity",
    "unit_price",
    "discount",
    "net_amount",
    "tax_rate",
    "tax_type",
    "tax_amount",
    "total_amount",
    "asin",
    "hsn",
    "invoice_tax_total",
    "invoice_total",
    "observed_total_amount",
    "total_resolution",
    "amount_in_words",
    "reverse_charge",
    "transaction_id",
    "payment_date_time",
    "invoice_value",
    "payment_mode",
    "validation",
]

print("\nCreating CSV file...")
with open(CSV_OUTPUT, "w", encoding="utf-8", newline="") as handle:
    writer = csv.DictWriter(handle, fieldnames=csv_fields, extrasaction="ignore")
    writer.writeheader()
    writer.writerows(all_rows)

print(f"CSV created: {CSV_OUTPUT}")


# ============================================================
# EXCEL WORKBOOK
# ============================================================

print("\nCreating Excel workbook...")
workbook = Workbook()

# ------------------------------------------------------------
# Sheet 1 — Invoice Details
# One row per detected document/page.
# ------------------------------------------------------------

details_sheet = workbook.active
details_sheet.title = "Invoice Details"

detail_fields = [
    "document_index",
    "document_type",
    "page",
    "invoice_number",
    "invoice_details",
    "invoice_date",
    "order_number",
    "order_date",
    "seller_name",
    "seller_address",
    "seller_pan",
    "seller_gstin",
    "billing_name",
    "billing_address",
    "billing_state_code",
    "shipping_name",
    "shipping_address",
    "shipping_state_code",
    "place_of_supply",
    "place_of_delivery",
    "invoice_tax_total",
    "invoice_total",
    "observed_total_amount",
    "total_resolution",
    "amount_in_words",
    "reverse_charge",
    "transaction_id",
    "payment_date_time",
    "invoice_value",
    "payment_mode",
]

details_sheet.append(detail_fields)
for _, context in contexts:
    details_sheet.append([context.get(field) for field in detail_fields])

# ------------------------------------------------------------
# Sheet 2 — Items & Charges
# Keeps both products and non-product charges.
# ------------------------------------------------------------

items_sheet = workbook.create_sheet("Items & Charges")
item_fields = [
    "document_index",
    "page",
    "invoice_number",
    "seller_name",
    "row_type",
    "charge_type",
    "description",
    "quantity",
    "unit_price",
    "discount",
    "net_amount",
    "tax_rate",
    "tax_type",
    "tax_amount",
    "total_amount",
    "asin",
    "hsn",
    "validation",
]
items_sheet.append(item_fields)

for row in all_rows:
    if row.get("row_type") == "document":
        continue
    items_sheet.append([row.get(field) for field in item_fields])

# ------------------------------------------------------------
# Sheet 3 — Products
# Product-only view for convenient downstream use.
# ------------------------------------------------------------

products_sheet = workbook.create_sheet("Products")
product_fields = [
    "document_index",
    "page",
    "invoice_number",
    "seller_name",
    "description",
    "quantity",
    "unit_price",
    "discount",
    "net_amount",
    "tax_rate",
    "tax_type",
    "tax_amount",
    "total_amount",
    "asin",
    "hsn",
    "validation",
]
products_sheet.append(product_fields)
for row in all_rows:
    if row.get("row_type") == "product":
        products_sheet.append([row.get(field) for field in product_fields])

# ------------------------------------------------------------
# Sheet 4 — Charges
# Charge/service/COD/shipping rows are kept separately.
# ------------------------------------------------------------

charges_sheet = workbook.create_sheet("Charges")
charge_fields = [
    "document_index",
    "page",
    "invoice_number",
    "seller_name",
    "charge_type",
    "description",
    "quantity",
    "unit_price",
    "discount",
    "net_amount",
    "tax_rate",
    "tax_type",
    "tax_amount",
    "total_amount",
    "validation",
]
charges_sheet.append(charge_fields)
for row in all_rows:
    if row.get("row_type") == "charge":
        charges_sheet.append([row.get(field) for field in charge_fields])

# ------------------------------------------------------------
# Sheet 5 — Validation Summary
# ------------------------------------------------------------

validation_sheet = workbook.create_sheet("Validation Summary")
validation_sheet.append(["Scope", "Check", "Status", "Message", "Severity"])

if validation:
    checks = as_list(validation.get("checks"))
    for check in checks:
        if not isinstance(check, dict):
            continue
        validation_sheet.append([
            check.get("scope") or check.get("document"),
            check.get("check"),
            check.get("status"),
            check.get("message"),
            check.get("severity"),
        ])

    summary = as_dict(validation.get("summary"))
    if summary:
        validation_sheet.append([])
        validation_sheet.append(["Summary", "Overall Status", validation.get("overall_status"), "", ""])
        validation_sheet.append(["Summary", "Total Checks", summary.get("total_checks"), "", ""])
        validation_sheet.append(["Summary", "Passed", summary.get("passed"), "", ""])
        validation_sheet.append(["Summary", "Failed", summary.get("failed"), "", ""])
        validation_sheet.append(["Summary", "Review", summary.get("review"), "", ""])
else:
    validation_sheet.append([
        "System",
        "Validation report",
        "UNAVAILABLE",
        "validation_report.json was not found.",
        "warning",
    ])


# ------------------------------------------------------------
# Sheet 6 — Export Summary
# ------------------------------------------------------------

summary_sheet = workbook.create_sheet("Export Summary")
summary_sheet.append(["Metric", "Value"])
summary_sheet.append(["Documents exported", len(contexts)])
summary_sheet.append(["Rows exported", len(all_rows)])
summary_sheet.append([
    "Products exported",
    sum(1 for row in all_rows if row.get("row_type") == "product"),
])
summary_sheet.append([
    "Charges exported",
    sum(1 for row in all_rows if row.get("row_type") == "charge"),
])
summary_sheet.append([
    "Other table rows preserved",
    sum(1 for row in all_rows if row.get("row_type") == "other"),
])


# ============================================================
# FORMATTING
# ============================================================

for sheet in workbook.worksheets:
    sheet.freeze_panes = "A2"

    for column_cells in sheet.columns:
        if not column_cells:
            continue

        max_length = 0
        column_index = column_cells[0].column
        column_letter = get_column_letter(column_index)

        for cell in column_cells:
            if cell.value is not None:
                max_length = max(max_length, len(str(cell.value)))

        sheet.column_dimensions[column_letter].width = min(
            max(max_length + 2, 12),
            60,
        )


workbook.save(EXCEL_OUTPUT)
print(f"Excel created: {EXCEL_OUTPUT}")


# ============================================================
# SUMMARY
# ============================================================

product_count = sum(1 for row in all_rows if row.get("row_type") == "product")
charge_count = sum(1 for row in all_rows if row.get("row_type") == "charge")
other_count = sum(1 for row in all_rows if row.get("row_type") == "other")

print("\n" + "=" * 70)
print("EXPORT COMPLETE")
print("=" * 70)
print(f"Documents exported : {len(contexts)}")
print(f"Rows exported      : {len(all_rows)}")
print(f"Products exported  : {product_count}")
print(f"Charges exported   : {charge_count}")
print(f"Other rows kept    : {other_count}")
print(f"\nCSV:   {CSV_OUTPUT}")
print(f"Excel: {EXCEL_OUTPUT}")
print("\nExcel sheets:")
print("  1. Invoice Details")
print("  2. Items & Charges")
print("  3. Products")
print("  4. Charges")
print("  5. Validation Summary")
print("  6. Export Summary")
print("=" * 70)
